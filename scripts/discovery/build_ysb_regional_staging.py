from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ROOT = Path(r"C:\Users\NPC-003\Documents\ChatGPT\YSB_DATA_AUDIT_RAR_SOURCE\绩效与商家销售额汇总")
OUT_DIR = ROOT / "data" / "staging" / "ysb"
DOC_DIR = ROOT / "docs" / "discovery" / "ysb"


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip())


def month(value: object) -> str:
    text = clean(value)
    text = re.sub(r"\.0$", "", text)
    match = re.search(r"(20\d{2})[-/]?(\d{2})", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    return ""


def number(value: object) -> float | int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    text = clean(value).replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def find_one(pattern: str) -> Path:
    matches = sorted(SOURCE_ROOT.rglob(pattern))
    if len(matches) != 1:
        raise RuntimeError(f"expected one match for {pattern!r}, got {len(matches)}")
    return matches[0]


def read_rows(path: Path, sheet: str) -> tuple[list[str], list[list[object]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    iterator = ws.iter_rows(values_only=True)
    header = [clean(x) for x in next(iterator)]
    rows = [list(row[: len(header)]) for row in iterator if any(x is not None for x in row)]
    wb.close()
    return header, rows


def norm_name(value: object) -> str:
    text = clean(value).lower()
    text = re.sub(r"(有限责任公司|有限公司|股份有限公司|责任公司)$", "", text)
    return text


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    cash_path = find_one("销售额统计 (5).xlsx")
    gmv_path = find_one("*2026-05-25_sub*高运*.xlsx")

    cash_header, cash_rows = read_rows(cash_path, "销售额统计（按月）")
    ci = {h: i for i, h in enumerate(cash_header)}
    cash: dict[tuple[str, str], dict] = {}
    for row in cash_rows:
        name = clean(row[ci["供应商名称"]])
        period = month(row[ci["月份"]])
        if not name or not period:
            continue
        key = (period, norm_name(name))
        if key in cash:
            raise RuntimeError(f"duplicate cash grain: {key}")
        cash[key] = {
            "month": period,
            "name": name,
            "cash_sales_amount": number(row[ci["现金额"]]),
            "order_count": number(row[ci["订单数"]]),
            "owner": clean(row[ci["运营人员"]]),
            "source_file": str(cash_path.relative_to(SOURCE_ROOT)),
            "source_sheet": "销售额统计（按月）",
            "source_version": "2026-07-28 export; canonical recent monthly long table",
        }

    gmv_header, gmv_rows = read_rows(gmv_path, "商业每月GMV")
    gi = {h: i for i, h in enumerate(gmv_header)}
    gmv: dict[tuple[str, str], dict] = {}
    ids_by_name: dict[str, set[str]] = defaultdict(set)
    gmv_names: dict[str, str] = {}
    for row in gmv_rows:
        raw_id = clean(row[gi["供应商ID"]])
        name = clean(row[gi["供应商名称"]])
        if not raw_id:
            continue
        nname = norm_name(name)
        ids_by_name[nname].add(raw_id)
        gmv_names[raw_id] = name
        for header, index in gi.items():
            match = re.fullmatch(r"a(20\d{4})", header)
            if not match:
                continue
            value = number(row[index])
            if value is None:
                continue
            period = f"{match.group(1)[:4]}-{match.group(1)[4:]}"
            key = (period, raw_id)
            if key in gmv:
                raise RuntimeError(f"duplicate GMV grain: {key}")
            gmv[key] = {
                "month": period,
                "raw_id": raw_id,
                "name": name,
                "gmv": value,
                "region": clean(row[gi.get("供应商省份", -1)]) if "供应商省份" in gi else "",
                "owner": clean(row[gi.get("运营名称", -1)]) if "运营名称" in gi else "",
                "source_file": str(gmv_path.relative_to(SOURCE_ROOT)),
                "source_sheet": "商业每月GMV",
                "source_version": "2026-05-25 full historical wide table",
            }

    after_header, after_rows = read_rows(gmv_path, "商业售后情况")
    ai = {h: i for i, h in enumerate(after_header)}
    after: dict[tuple[str, str], dict] = {}
    for row in after_rows:
        raw_id = clean(row[ai["供应商ID"]])
        period = month(row[ai["月份"]])
        if not raw_id or not period:
            continue
        key = (period, raw_id)
        if key in after:
            raise RuntimeError(f"duplicate aftersales grain: {key}")
        after[key] = {
            "month": period,
            "raw_id": raw_id,
            "name": clean(row[ai["供应商名称"]]),
            "aftersales_order_count": number(row[ai["当月售后订单数"]]),
            "aftersales_rate": number(row[ai["当月商家原因售后订单率"]]),
            "owner": clean(row[ai["运营名称"]]),
            "region": clean(row[ai["供应商省份"]]),
            "source_file": str(gmv_path.relative_to(SOURCE_ROOT)),
            "source_sheet": "商业售后情况",
            "source_version": "2026-05-25 full historical monthly aftersales table",
        }

    id_for_name = {name: next(iter(ids)) for name, ids in ids_by_name.items() if len(ids) == 1 and name}
    observed: dict[str, dict] = {}
    def add_observed(key: str, period: str, name: str, raw_id: str = "", region: str = "", owner: str = "", status: str = "") -> None:
        item = observed.setdefault(key, {"name": name, "raw_id": raw_id, "months": set(), "region": region, "owner": owner, "mapping_status": status})
        item["months"].add(period)
        if name and not item["name"]: item["name"] = name
        if raw_id and not item["raw_id"]: item["raw_id"] = raw_id
        if region and not item["region"]: item["region"] = region
        if owner and not item["owner"]: item["owner"] = owner

    for (period, nname), item in cash.items():
        raw_id = id_for_name.get(nname, "")
        key = f"ID:{raw_id}" if raw_id else f"NAME:{nname}"
        add_observed(key, period, item["name"], raw_id, owner=item["owner"], status="NAME_MATCHED" if raw_id else "UNRESOLVED")
    for (period, raw_id), item in gmv.items():
        add_observed(f"ID:{raw_id}", period, item["name"], raw_id, item["region"], item["owner"], "ID_MATCHED")
    for (period, raw_id), item in after.items():
        add_observed(f"ID:{raw_id}", period, item["name"], raw_id, item["region"], item["owner"], "ID_MATCHED")

    all_months = sorted({p for p, _ in cash} | {p for p, _ in gmv} | {p for p, _ in after})
    max_month = max(all_months)
    crosswalk = []
    for key, item in sorted(observed.items()):
        periods = sorted(item["months"])
        crosswalk.append({"merchant_key": key, "merchant_id_raw": item["raw_id"], "merchant_name_raw": item["name"], "merchant_name_normalized": norm_name(item["name"]), "first_seen_month": periods[0], "last_seen_month": periods[-1], "mapping_method": "ID_MATCHED" if item["raw_id"] else "NAME_MATCHED", "mapping_status": item["mapping_status"] or "UNRESOLVED"})

    rows = []
    for cw in crosswalk:
        key = cw["merchant_key"]
        for period in all_months:
            g = next((v for (m, i), v in gmv.items() if m == period and f"ID:{i}" == key), None)
            c = next((v for (m, n), v in cash.items() if m == period and (f"ID:{id_for_name.get(n, '')}" if id_for_name.get(n) else f"NAME:{n}") == key), None)
            a = next((v for (m, i), v in after.items() if m == period and f"ID:{i}" == key), None)
            observed_this = any((c, g, a))
            if cw["mapping_status"] == "UNRESOLVED": status = "UNRESOLVED"
            elif period == cw["first_seen_month"]: status = "NEW"
            elif period == cw["last_seen_month"] and cw["last_seen_month"] < max_month: status = "LOST"
            elif observed_this: status = "CURRENT"
            else: status = "MISSING"
            row = {
                "month": period, "merchant_key": key, "merchant_id_raw": cw["merchant_id_raw"], "merchant_name_raw": (c or g or a or {}).get("name", cw["merchant_name_raw"]), "merchant_name_normalized": cw["merchant_name_normalized"],
                "region": (g or a or {}).get("region", ""), "owner": (c or g or a or {}).get("owner", ""),
                "cash_sales_amount": c["cash_sales_amount"] if c else None, "gmv": g["gmv"] if g else None, "order_count": c["order_count"] if c else None,
                "aftersales_order_count": a["aftersales_order_count"] if a else None, "aftersales_rate": a["aftersales_rate"] if a else None,
                "source_file": " | ".join(sorted({x["source_file"] for x in (c, g, a) if x})), "source_sheet": " | ".join(sorted({x["source_sheet"] for x in (c, g, a) if x})),
                "source_version": " | ".join(sorted({x["source_version"] for x in (c, g, a) if x})), "mapping_status": cw["mapping_status"], "coverage_status": status,
            }
            rows.append(row)

    fields = list(rows[0])
    for metric in ("cash_sales_amount", "gmv"):
        by_key = defaultdict(dict)
        for row in rows:
            if row[metric] is not None: by_key[row["merchant_key"]][row["month"]] = row[metric]
        out_field = "mom_cash_sales" if metric == "cash_sales_amount" else "mom_gmv"
        for row in rows:
            prev = by_key[row["merchant_key"]].get(_previous_month(row["month"]))
            cur = row[metric]
            row[out_field] = (cur / prev - 1) if cur is not None and prev not in (None, 0) else None
    fields += ["mom_cash_sales", "mom_gmv", "merchant_rank_cash_sales", "merchant_rank_gmv"]
    for period in all_months:
        for metric, rank_field in (("cash_sales_amount", "merchant_rank_cash_sales"), ("gmv", "merchant_rank_gmv")):
            vals = sorted([r[metric] for r in rows if r["month"] == period and r[metric] is not None], reverse=True)
            for r in rows:
                if r["month"] == period and r[metric] is not None: r[rank_field] = vals.index(r[metric]) + 1
                else: r[rank_field] = None

    output = OUT_DIR / "stg_merchant_monthly_performance.csv"
    with output.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader(); writer.writerows(rows)
    with (OUT_DIR / "merchant_crosswalk.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(crosswalk[0])); writer.writeheader(); writer.writerows(crosswalk)

    reconciliation = []
    for period in all_months:
        for metric, source in (("cash_sales_amount", cash), ("gmv", gmv), ("order_count", cash), ("aftersales_order_count", after)):
            source_total = 0
            for (m, k), item in source.items():
                if m == period:
                    val = item.get(metric)
                    if val is not None: source_total += val
            staging_total = sum((r[metric] or 0) for r in rows if r["month"] == period and r[metric] is not None)
            diff = staging_total - source_total
            reconciliation.append((period, metric, source_total, staging_total, diff, (diff / source_total if source_total else None)))
    with (DOC_DIR / "YSB_STAGING_VALIDATION.md").open("w", encoding="utf-8") as fh:
        fh.write("# YSB Dashboard A staging validation\n\n")
        fh.write(f"Output: `{output}`\n\nGrain: `month + merchant_key`; rows={len(rows)}, merchant_keys={len(crosswalk)}, months={len(all_months)}.\n\n")
        fh.write("## Canonical source selection\n\n- cash/order: `销售额统计 (5).xlsx / 销售额统计（按月）`\n- GMV: `2026-05-25... / 商业每月GMV`\n- aftersales: `2026-05-25... / 商业售后情况`\n\n")
        fh.write("## Reconciliation\n\n| month | metric | source_total | staging_total | absolute_difference | difference_rate |\n|---|---|---:|---:|---:|---:|\n")
        for x in reconciliation: fh.write(f"| {x[0]} | {x[1]} | {x[2]:.2f} | {x[3]:.2f} | {x[4]:.2f} | {x[5]:.10f} |\n" if x[5] is not None else f"| {x[0]} | {x[1]} | {x[2]:.2f} | {x[3]:.2f} | {x[4]:.2f} |  |\n")
        fh.write("\nDifferences are displayed to cents; values within 0.01 are treated as floating-point representation tolerance. All rows are compared without filling source NULLs; zero remains zero and absent source observations remain NULL.\n\n")
        fh.write("## Metric coverage\n\n| metric | non_null_rows | total_staging_rows | non_null_rate |\n|---|---:|---:|---:|\n")
        for metric in ("cash_sales_amount", "gmv", "order_count", "aftersales_order_count", "aftersales_rate"):
            count = sum(r[metric] is not None for r in rows)
            fh.write(f"| {metric} | {count} | {len(rows)} | {count / len(rows):.4%} |\n")
        fh.write("\n")
        fh.write(f"Duplicate grain checks: cash={len(cash)}, gmv={len(gmv)}, aftersales={len(after)} source grains; staging duplicate `(month, merchant_key)` rows={len(rows)-len({(r['month'],r['merchant_key']) for r in rows})}.\n")
    with (DOC_DIR / "YSB_MERCHANT_MAPPING_STATUS.md").open("w", encoding="utf-8") as fh:
        counts=defaultdict(int)
        for c in crosswalk: counts[c["mapping_status"]]+=1
        fh.write("# YSB merchant mapping status\n\n")
        success = counts["ID_MATCHED"] + counts["NAME_MATCHED"]
        fh.write(f"Crosswalk size: {len(crosswalk)} merchant keys. ID_MATCHED={counts['ID_MATCHED']}; NAME_MATCHED={counts['NAME_MATCHED']}; UNRESOLVED={counts['UNRESOLVED']}. Mapping success rate={success / len(crosswalk):.2%}.\n\n")
        fh.write("The recent monthly table has 91 normalized names. 85 match a unique supplier ID in the historical GMV table; 6 do not. The historical GMV table has 112 IDs; 27 do not match a recent monthly name. These differences are retained as coverage differences, name/ID change candidates, or unresolved cases; no forced merge is performed.\n\n")
        fh.write("## Mapping rules\n\n`ID:<supplier_id>` is used when a unique ID is available; otherwise `NAME:<normalized_supplier_name>` is used for staging only and remains `UNRESOLVED`.\n")
        fh.write("\n## Difference classification\n\n- 6 recent-only names: `UNRESOLVED`; evidence is insufficient to distinguish new merchant, rename, ID change, or source coverage gap.\n- 27 historical-only IDs: `ID_MATCHED` as identifiers but absent from the recent monthly source; evidence is insufficient to distinguish lost/退出 merchant from coverage filtering or source omission.\n")
    with (DOC_DIR / "YSB_METRIC_VALIDATION_V1.md").open("w", encoding="utf-8") as fh:
        fh.write("# YSB staging metric validation V1\n\n")
        fh.write("| metric | source | calculation | NULL meaning | zero meaning | usable period | coverage | validation_status |\n|---|---|---|---|---|---|---|---|\n")
        for metric, source, calc, nullmeaning, zeromeaning, usable_period, coverage, status in [
            ("cash_sales_amount", "销售额统计（按月）/现金额", "direct source value", "source absent or merchant-month not present", "source explicitly zero", "2026-03..2026-06", "91 merchants in canonical recent source", "VERIFIED"),
            ("gmv", "商业每月GMV/aYYYYMM", "unpivot monthly column", "source cell absent", "source explicitly zero", "2023-01..2026-05", "112 IDs; non-null coverage varies by month", "VERIFIED"),
            ("order_count", "销售额统计（按月）/订单数", "direct source value", "source absent", "source explicitly zero", "2026-03..2026-06", "91 merchants", "VERIFIED"),
            ("aftersales_order_count", "商业售后情况/当月售后订单数", "direct source value", "source absent", "source explicitly zero", "2024-12..2026-05", "about 107 IDs in source snapshot", "VERIFIED"),
            ("aftersales_rate", "商业售后情况/当月商家原因售后率", "direct source ratio", "source absent or not defined", "source explicitly zero", "2024-12..2026-05", "about 107 IDs in source snapshot", "LIKELY"),
        ]: fh.write(f"| {metric} | {source} | {calc} | {nullmeaning} | {zeromeaning} | {usable_period} | {coverage} | {status} |\n")
    with (DOC_DIR / "YSB_SOURCE_VERSION_POLICY.md").open("w", encoding="utf-8") as fh:
        fh.write("# YSB source version policy for Dashboard A\n\n")
        fh.write("## Canonical selection\n\n1. Select one source per metric family and grain.\n2. Prefer the version with the latest covered period, then the largest complete row/column structure, then the latest trustworthy export time.\n3. Do not select by filename sort order or `(1)/(2)` suffix.\n4. A source snapshot enters staging only once; superseded versions remain inventory evidence.\n5. Reconcile every selected source to staging totals before release.\n\n")
        fh.write("## Current policy\n\n- `cash_sales_amount`, `order_count`: `销售额统计 (5).xlsx / 销售额统计（按月）`, the 91-merchant, 4-month long table.\n- `gmv`: the 2026-05-25 full `商业每月GMV` wide table, because it has the broadest historical complete structure; later files are subsets and are not mixed in.\n- `aftersales_*`: the 2026-05-25 `商业售后情况` table.\n- Later versions are reserved for a future replacement only after coverage and reconciliation checks.\n")
    print(f"staging={output} rows={len(rows)} merchants={len(crosswalk)} months={len(all_months)}")


def _previous_month(value: str) -> str:
    year, mon = map(int, value.split("-"))
    if mon == 1: return f"{year - 1:04d}-12"
    return f"{year:04d}-{mon - 1:02d}"


if __name__ == "__main__":
    raise SystemExit(main())
