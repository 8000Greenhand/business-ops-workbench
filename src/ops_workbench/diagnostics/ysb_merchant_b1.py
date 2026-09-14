"""Single-merchant order-item staging and monthly diagnostic marts."""

from __future__ import annotations

import calendar
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook

POLICY_USABLE_WITH_LIMITATION = "USABLE_WITH_LIMITATION"
STATUS_RETAINED = "RETAINED"
STATUS_CURRENT_ONLY = "CURRENT_ONLY"
STATUS_PREVIOUS_ONLY = "PREVIOUS_ONLY"
PRODUCT_STATUS_RETAINED = "RETAINED_ACTIVE"
PRODUCT_STATUS_CURRENT_ONLY = "CURRENT_ONLY_ACTIVE"
PRODUCT_STATUS_PREVIOUS_ONLY = "PREVIOUS_ONLY_ACTIVE"

SOURCE_COLUMNS = {
    "下单时间": "order_date",
    "订单ID": "order_id",
    "订单编号": "order_number",
    "药店编码": "customer_key",
    "药店全称": "customer_name",
    "商品ID": "product_key",
    "商品编码": "product_code",
    "产品名称": "product_name",
    "厂家": "manufacturer_raw",
    "规格": "specification_raw",
    "采购量": "quantity",
    "进货价": "unit_purchase_price",
    "进货金额": "purchase_amount",
    "订单状态": "order_status_raw",
    "处理情况": "processing_status_raw",
    "备注": "processing_remark_raw",
    "结算状态": "settlement_status_raw",
    "订单类型": "order_type_raw",
}


def load_order_metric_policy(path: Path) -> dict[str, Any]:
    """Load and minimally validate the explicit order metric policy."""
    policy = yaml.safe_load(path.read_text(encoding="utf-8"))
    required = {"name", "status", "purchase_amount", "order_count", "zero_amount", "observed_values"}
    missing = required.difference(policy)
    if missing:
        raise ValueError(f"order metric policy missing keys: {sorted(missing)}")
    if policy["purchase_amount"]["inclusion"] != "numeric_values_all_statuses":
        raise ValueError("unsupported purchase_amount inclusion policy")
    if policy["order_count"]["inclusion"] != "grouped_purchase_amount_gt_zero":
        raise ValueError("unsupported order_count inclusion policy")
    return policy


def read_order_items(path: Path, sheet_name: str = "订单原始数据") -> pd.DataFrame:
    """Read only the required order-item fields from the merchant workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [_clean_text(value) for value in next(iterator)]
        indexes = {name: headers.index(name) for name in SOURCE_COLUMNS}
        records: list[dict[str, object]] = []
        for source_row, row in enumerate(iterator, start=2):
            record = {
                target: _clean_text(row[indexes[source]])
                for source, target in SOURCE_COLUMNS.items()
            }
            record["source_row"] = source_row
            records.append(record)
    finally:
        workbook.close()

    frame = pd.DataFrame.from_records(records)
    frame["order_date"] = pd.to_datetime(frame["order_date"], errors="coerce").dt.normalize()
    for column in ("order_id", "order_number", "customer_key", "product_key", "product_code"):
        frame[column] = frame[column].map(_identifier).astype("string")
    for column in ("customer_name", "product_name", "manufacturer_raw", "specification_raw"):
        frame[column] = frame[column].astype("string")
    for column in ("order_status_raw", "processing_status_raw", "processing_remark_raw", "settlement_status_raw", "order_type_raw"):
        frame[column] = frame[column].replace("", pd.NA).astype("string")
    for column in ("quantity", "unit_purchase_price", "purchase_amount"):
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame["source_sheet"] = sheet_name
    return frame


def apply_order_policy(frame: pd.DataFrame, policy: dict[str, Any]) -> pd.DataFrame:
    """Apply the file-internal policy and add explicit data-quality flags."""
    result = frame.copy()
    tolerance = float(policy.get("amount_tolerance", 0.011))
    calculated = result["quantity"] * result["unit_purchase_price"]
    result["dq_missing_order_id"] = result["order_id"].isna() | result["order_id"].eq("")
    result["dq_missing_customer_key"] = result["customer_key"].isna() | result["customer_key"].eq("")
    result["dq_missing_product_key"] = result["product_key"].isna() | result["product_key"].eq("")
    result["dq_missing_purchase_amount"] = result["purchase_amount"].isna()
    result["dq_purchase_amount_mismatch"] = (
        result[["quantity", "unit_purchase_price", "purchase_amount"]].notna().all(axis=1)
        & (calculated - result["purchase_amount"]).abs().gt(tolerance)
    )
    result["dq_zero_amount_line"] = result["purchase_amount"].eq(0)

    customer_names = result.dropna(subset=["customer_key", "customer_name"]).groupby("customer_key")["customer_name"].nunique()
    customer_conflicts = set(customer_names[customer_names.gt(1)].index)
    result["dq_customer_name_conflict"] = result["customer_key"].isin(customer_conflicts)

    mapping = result.dropna(subset=["product_key"]).copy()
    mapping["product_signature"] = (
        mapping["manufacturer_raw"].fillna("") + "|" + mapping["specification_raw"].fillna("")
    )
    code_counts = mapping.groupby("product_key")["product_code"].nunique()
    signature_counts = mapping.groupby("product_key")["product_signature"].nunique()
    product_conflicts = set(code_counts[code_counts.gt(1)].index) | set(signature_counts[signature_counts.gt(1)].index)
    result["dq_product_mapping_conflict"] = result["product_key"].isin(product_conflicts)

    dates_per_order = result.dropna(subset=["order_id", "order_date"]).groupby("order_id")["order_date"].nunique()
    cross_date_orders = set(dates_per_order[dates_per_order.gt(1)].index)
    result["dq_order_cross_date"] = result["order_id"].isin(cross_date_orders)

    included_amount = result["purchase_amount"].fillna(0.0)
    result["purchase_amount_included"] = result["purchase_amount"].notna()
    result["policy_purchase_amount"] = included_amount
    order_amount = result.groupby("order_id", dropna=False)["policy_purchase_amount"].transform("sum")
    result["order_purchase_amount"] = order_amount
    result["included_in_order_count"] = result["order_id"].notna() & order_amount.gt(0)
    result["dq_zero_amount_order"] = result["order_id"].notna() & order_amount.eq(0)

    observed = policy["observed_values"]
    for column in ("order_status_raw", "processing_status_raw", "settlement_status_raw"):
        result[f"dq_unrecognized_{column}"] = ~result[column].fillna("<BLANK>").isin(observed.get(column, []))
    result["status_semantics_quality"] = "UNCONFIRMED"
    result["metric_policy_name"] = policy["name"]
    result["metric_policy_status"] = policy["status"]

    flag_columns = [column for column in result.columns if column.startswith("dq_")]
    result["data_quality_flags"] = result.apply(
        lambda row: ";".join(column.removeprefix("dq_").upper() for column in flag_columns if bool(row[column])),
        axis=1,
    )
    return result


def build_monthly_diagnosis(items: pd.DataFrame) -> pd.DataFrame:
    """Build monthly purchase metrics and changes for adjacent complete months."""
    frame = items.copy()
    frame["month"] = frame["order_date"].dt.to_period("M")
    rows: list[dict[str, object]] = []
    for month, group in frame.dropna(subset=["month"]).groupby("month", observed=True):
        days = group["order_date"].dt.day.nunique()
        calendar_days = calendar.monthrange(month.year, month.month)[1]
        complete = group["order_date"].dt.day.min() == 1 and group["order_date"].dt.day.max() == calendar_days and days == calendar_days
        positive = group["policy_purchase_amount"].gt(0)
        order_count = group.loc[group["included_in_order_count"], "order_id"].nunique()
        purchase_amount = group["policy_purchase_amount"].sum()
        rows.append(
            {
                "month": str(month),
                "purchase_amount": purchase_amount,
                "order_count": order_count,
                "aov": purchase_amount / order_count if order_count else pd.NA,
                "active_customers": group.loc[positive, "customer_key"].nunique(),
                "active_products": group.loc[positive, "product_key"].nunique(),
                "coverage_start": group["order_date"].min().date().isoformat(),
                "coverage_end": group["order_date"].max().date().isoformat(),
                "observed_days": days,
                "calendar_days": calendar_days,
                "period_quality_status": "COMPLETE" if complete else "PARTIAL",
                "metric_policy_name": group["metric_policy_name"].iloc[0],
                "metric_policy_status": group["metric_policy_status"].iloc[0],
            }
        )
    result = pd.DataFrame(rows).sort_values("month").reset_index(drop=True)
    metrics = ("purchase_amount", "order_count", "aov", "active_customers", "active_products")
    for metric in metrics:
        result[f"{metric}_change"] = pd.NA
        result[f"{metric}_change_rate"] = pd.NA
    result["previous_month"] = pd.NA
    result["comparison_quality_status"] = "UNAVAILABLE"
    by_month = result.set_index("month")
    for index, row in result.iterrows():
        current = pd.Period(row["month"], freq="M")
        previous = str(current - 1)
        if previous not in by_month.index:
            continue
        prior = by_month.loc[previous]
        if row["period_quality_status"] != "COMPLETE" or prior["period_quality_status"] != "COMPLETE":
            result.at[index, "comparison_quality_status"] = "INCOMPLETE_PERIOD"
            continue
        result.at[index, "previous_month"] = previous
        result.at[index, "comparison_quality_status"] = "VERIFIED"
        for metric in metrics:
            current_value = row[metric]
            previous_value = prior[metric]
            result.at[index, f"{metric}_change"] = current_value - previous_value
            result.at[index, f"{metric}_change_rate"] = (
                current_value / previous_value - 1 if pd.notna(previous_value) and previous_value != 0 else pd.NA
            )
    return result


def latest_complete_pair(monthly: pd.DataFrame) -> tuple[pd.Period, pd.Period]:
    """Return the latest adjacent pair of complete calendar months."""
    complete = {pd.Period(value, freq="M") for value in monthly.loc[monthly["period_quality_status"].eq("COMPLETE"), "month"]}
    for current in sorted(complete, reverse=True):
        if current - 1 in complete:
            return current - 1, current
    raise ValueError("no adjacent complete month pair")


def build_customer_contribution(items: pd.DataFrame, previous: pd.Period, current: pd.Period) -> pd.DataFrame:
    """Build objective customer purchase-amount changes for two periods."""
    return _build_contribution(items, previous, current, "customer_key", "customer_name", False)


def build_product_contribution(items: pd.DataFrame, previous: pd.Period, current: pd.Period) -> pd.DataFrame:
    """Build objective product purchase-amount changes and mapping quality."""
    result = _build_contribution(items, previous, current, "product_key", "product_name", True)
    source = items.dropna(subset=["product_key"]).copy()
    source["signature"] = source["manufacturer_raw"].fillna("") + "|" + source["specification_raw"].fillna("")
    codes = source.groupby("product_key")["product_code"].nunique()
    signatures = source.groupby("product_key")["signature"].nunique()
    result["mapping_quality_status"] = result["product_key"].map(
        lambda key: "CONFLICT" if codes.get(key, 0) > 1 or signatures.get(key, 0) > 1 else "VALID"
    )
    result["product_code"] = result["product_key"].map(_preferred_value(source, "product_key", "product_code"))
    result["manufacturer"] = result["product_key"].map(_preferred_value(source, "product_key", "manufacturer_raw"))
    result["specification"] = result["product_key"].map(_preferred_value(source, "product_key", "specification_raw"))
    return result[
        [
            "previous_month", "current_month", "product_key", "product_code", "product_name",
            "manufacturer", "specification", "previous_amount", "current_amount", "amount_change",
            "period_status", "mapping_quality_status",
        ]
    ]


def build_quality_summary(items: pd.DataFrame) -> pd.DataFrame:
    """Summarize row and entity counts for every explicit data-quality rule."""
    rules = [
        ("MISSING_ORDER_ID", "dq_missing_order_id", "order_id"),
        ("MISSING_CUSTOMER_KEY", "dq_missing_customer_key", "customer_key"),
        ("CUSTOMER_KEY_NAME_CONFLICT", "dq_customer_name_conflict", "customer_key"),
        ("MISSING_PRODUCT_KEY", "dq_missing_product_key", "product_key"),
        ("PRODUCT_KEY_MAPPING_CONFLICT", "dq_product_mapping_conflict", "product_key"),
        ("ZERO_AMOUNT_LINE", "dq_zero_amount_line", "source_row"),
        ("ZERO_AMOUNT_ORDER", "dq_zero_amount_order", "order_id"),
        ("UNRECOGNIZED_ORDER_STATUS", "dq_unrecognized_order_status_raw", "order_status_raw"),
        ("UNRECOGNIZED_PROCESSING_STATUS", "dq_unrecognized_processing_status_raw", "processing_status_raw"),
        ("UNRECOGNIZED_SETTLEMENT_STATUS", "dq_unrecognized_settlement_status_raw", "settlement_status_raw"),
        ("ORDER_CROSS_DATE", "dq_order_cross_date", "order_id"),
        ("MISSING_PURCHASE_AMOUNT", "dq_missing_purchase_amount", "source_row"),
        ("PURCHASE_AMOUNT_MISMATCH", "dq_purchase_amount_mismatch", "source_row"),
    ]
    rows = []
    for rule, flag, entity in rules:
        affected = items[items[flag]]
        rows.append(
            {
                "rule": rule,
                "affected_rows": len(affected),
                "affected_entities": affected[entity].nunique(dropna=True),
                "quality_status": "VALID" if affected.empty else "REVIEW_REQUIRED",
            }
        )
    rows.append(
        {
            "rule": "UNCONFIRMED_STATUS_SEMANTICS",
            "affected_rows": int(items["status_semantics_quality"].eq("UNCONFIRMED").sum()),
            "affected_entities": items["order_id"].nunique(),
            "quality_status": POLICY_USABLE_WITH_LIMITATION,
        }
    )
    return pd.DataFrame(rows)


def observed_status_distribution(items: pd.DataFrame) -> pd.DataFrame:
    """Return source status values and counts without assigning business meaning."""
    records = []
    for column in ("order_status_raw", "processing_status_raw", "processing_remark_raw", "settlement_status_raw", "order_type_raw"):
        counts = Counter(items[column].fillna("<BLANK>").astype(str))
        records.extend({"field": column, "value": value, "row_count": count} for value, count in counts.most_common())
    return pd.DataFrame(records)


def _build_contribution(
    items: pd.DataFrame,
    previous: pd.Period,
    current: pd.Period,
    key: str,
    label: str,
    product: bool,
) -> pd.DataFrame:
    frame = items.copy()
    frame["month"] = frame["order_date"].dt.to_period("M")
    active = frame[frame["policy_purchase_amount"].gt(0) & frame["month"].isin([previous, current])]
    amounts = active.groupby(["month", key], observed=True)["policy_purchase_amount"].sum().unstack(level=0, fill_value=0)
    for period in (previous, current):
        if period not in amounts.columns:
            amounts[period] = 0.0
    amounts = amounts[[previous, current]]
    preferred = _preferred_value(active, key, label)
    result = amounts.reset_index().rename(columns={previous: "previous_amount", current: "current_amount"})
    result[label] = result[key].map(preferred)
    result["previous_month"] = str(previous)
    result["current_month"] = str(current)
    result["amount_change"] = result["current_amount"] - result["previous_amount"]
    result["period_status"] = result.apply(
        lambda row: _period_status(row["previous_amount"], row["current_amount"], product), axis=1
    )
    columns = ["previous_month", "current_month", key, label, "previous_amount", "current_amount", "amount_change", "period_status"]
    return result[columns].sort_values(["amount_change", key]).reset_index(drop=True)


def _period_status(previous: float, current: float, product: bool) -> str:
    if previous > 0 and current > 0:
        return PRODUCT_STATUS_RETAINED if product else STATUS_RETAINED
    if current > 0:
        return PRODUCT_STATUS_CURRENT_ONLY if product else STATUS_CURRENT_ONLY
    return PRODUCT_STATUS_PREVIOUS_ONLY if product else STATUS_PREVIOUS_ONLY


def _preferred_value(frame: pd.DataFrame, key: str, value: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for entity, group in frame.dropna(subset=[key]).groupby(key, observed=True):
        values = group[value].dropna().astype(str)
        if not values.empty:
            result[str(entity)] = sorted(values.value_counts().items(), key=lambda item: (-item[1], item[0]))[0][0]
    return result


def _clean_text(value: object) -> object:
    return value.strip() if isinstance(value, str) else value


def _identifier(value: object) -> object:
    if value is None or value == "":
        return pd.NA
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
